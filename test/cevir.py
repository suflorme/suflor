#!/usr/bin/env python3
# Proje hazırlığı T2 — metne çevirme (`baglam.py cevir`, cevir-isci.py): yapay klasörde görüntü, taranmış PDF, sunum, e-posta metne çevrilip
# dizinde aranabiliyor; içerikle gizli bilgi (şifre görüntüsü, API anahtarı, kart, IBAN) dizine girmiyor; adla hariç tutulan ve birebir
# kopya çevrilmiyor; önbellek (ikinci tur 0 çevirme); takılan dosyada zaman aşımı + yeni işçi; toplantı başlayınca durma ve kaldığı yerden.
# Gerçek Docling + Vision ister (<ortak>/cevir-venv; yoksa atlanır). ~1 dk.
#   PYTHONDONTWRITEBYTECODE=1 python3 test/cevir.py      (kaldı → çıkış 1)
import importlib.util, json, os, shutil, subprocess, sys, tempfile, time
sys.dont_write_bytecode = True; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__))); import bosport  # noqa: canlı kullanım koruması (gerçek model)

KOD = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ORTAK = "/Users/Shared/Suflor"; VPY = os.path.join(ORTAK, "cevir-venv", "bin", "python")
if not os.path.exists(VPY): print(f"çevirme: atlandı (ortam yok: {VPY})"); sys.exit(0)
T = tempfile.mkdtemp(prefix="suflor-cevir-"); P = os.path.join(T, "proje"); U = os.path.join(T, "uyg"); os.makedirs(P); os.makedirs(os.path.join(U, "canli"))
URET = r'''
import sys, os
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
P = sys.argv[1]; F = "/System/Library/Fonts/Supplemental/Arial.ttf"
def gorsel(satirlar, boy=34):
    im = Image.new("RGB", (1400, 120 + 70 * len(satirlar)), "white"); d = ImageDraw.Draw(im); f = ImageFont.truetype(F, boy)
    for i, s in enumerate(satirlar): d.text((60, 60 + 70 * i), s, fill="black", font=f)
    return im
os.makedirs(f"{P}/ekran", exist_ok=True); os.makedirs(f"{P}/belgeler", exist_ok=True); os.makedirs(f"{P}/eski", exist_ok=True)
gorsel(["Depo sayımı Çarşamba günü yapılacak", "Sorumlu: Gülşen Hanım, Bölge: İzmir"]).save(f"{P}/ekran/depo-sayimi.png")
gorsel(["Muhasebe portalı giriş bilgisi", "Kullanıcı: muhasebe01", "Şifre: Kış2024!Portal"]).save(f"{P}/ekran/giris-bilgisi.png")
gorsel(["Taranmış sözleşme eki: teslimat Ağustos ayında,", "cezai şart yüzde onbeş, Kocaeli fabrikası"]).save(f"{P}/belgeler/taranmis-sozlesme.pdf")
gorsel(["Kimlik kartı örneği"]).save(f"{P}/belgeler/kimlik-on.jpg")  # adla hariç: çevrilmez
gorsel(["Depo sayımı Çarşamba günü yapılacak", "Sorumlu: Gülşen Hanım, Bölge: İzmir"]).save(f"{P}/eski/depo-sayimi.png")  # birebir kopya
pr = Presentation()
for b, m in (("Yol haritası", "Birinci çeyrekte Ankara şubesi açılıyor"), ("Bütçe", "Pazarlama bütçesi yüzde yirmi artırıldı")):
    s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = b; s.placeholders[1].text = m
pr.save(f"{P}/belgeler/sunum.pptx")
'''
subprocess.run([VPY, "-B", "-c", URET, P], check=True, capture_output=True)
for k, m in (("posta/teklif.eml", "Subject: Teklif revizyonu\nFrom: Tedarik Birimi\nTo: Satin Alma\nDate: Mon, 12 Oct 2026 10:00:00 +0300\n"
                                    "Content-Type: text/plain; charset=utf-8\n\nYeni teklifte Erzurum teslimatı iki hafta öne çekildi.\n"),
             ("notlar/anahtar.md", "# Entegrasyon\n\nÖdeme servisinin hesap bağlantısı kuruldu.\n\nAPI: sk-ant-api03-AbCdEf0123456789ghIJklmnOPqrstUVwx kullanılıyor.\n\n"
                                   "## Plan\n\nRize lojistik merkezi haftaya devreye giriyor.\n"),
             ("notlar/odeme.md", "# Ödeme\n\nKart 4539 1488 0343 6467 ile ödendi.\n\n## Banka\n\nIBAN TR33 0006 1005 1978 6457 8413 26 hesabına.\n\n## Ek\n\nŞifreleme politikası güncellendi, parola en az 12 karakter.\n")):
    y = os.path.join(P, k); os.makedirs(os.path.dirname(y), exist_ok=True); open(y, "w", encoding="utf-8").write(m)
ay = os.path.join(T, "ayar.json"); json.dump({"alan": "Deneme", "uygulama": U, "proje": P, "ortak": ORTAK}, open(ay, "w"))
os.environ["SUFLOR_AYAR"] = ay; sys.dont_write_bytecode = True
def yukle():
    sp = importlib.util.spec_from_file_location("baglam_cevir", os.path.join(KOD, "baglam.py")); m = importlib.util.module_from_spec(sp); sp.loader.exec_module(m); return m
m = yukle(); m.CEVIR_ISCI = os.path.join(KOD, "cevir-isci.py")
hatalar = []
def kontrol(ad, kosul, ek=""):
    print(("✓ " if kosul else "✗ ") + ad + (f" — {ek}" if ek and not kosul else ""))
    if not kosul: hatalar.append(ad)
bul = lambda c, q: {r[1] for r in m.ara(q, n=10, guncel=False)[0]}
c = m.baglan()
try:
    m.guncelle(c)
    kontrol("çevrilmeden önce görüntü/PDF/sunum dizinde yok", not c.execute("SELECT 1 FROM dosya WHERE yol LIKE '%.png' OR yol LIKE '%.pdf' OR yol LIKE '%.pptx'").fetchone())
    kontrol("e-posta (.eml) doğrudan dizinde", "posta/teklif.eml" in bul(c, "Erzurum teslimatı"))
    r = m.cevir(c, sessiz=True); print("  " + m.cevir_ozet(r).replace("\n", "\n  "))
    kontrol("çevrildi: 2 görüntü + taranmış PDF + sunum (kimlik adla hariç, kopya atlandı)", r["cevrildi"] == 4 and r["hata"] == 0 and r["aday"] == 4, json.dumps(r, ensure_ascii=False)[:300])
    kontrol("görüntüdeki Türkçe metin aranıyor", "ekran/depo-sayimi.png" in bul(c, "depo sayımı Gülşen"))
    kontrol("taranmış PDF'teki metin aranıyor", "belgeler/taranmis-sozlesme.pdf" in bul(c, "cezai şart Kocaeli"))
    sl = c.execute("SELECT konum FROM parca WHERE yol='belgeler/sunum.pptx'").fetchall()
    kontrol("sunum slayt slayt, konumda slayt numarası", "belgeler/sunum.pptx" in bul(c, "Ankara şubesi") and any(k[0].startswith("slayt 2") for k in sl), str(sl))
    kontrol("şifreli ekran görüntüsünden parça dizine girmedi", not bul(c, "muhasebe portalı") and c.execute("SELECT gizli_parca FROM dosya WHERE yol='ekran/giris-bilgisi.png'").fetchone() == (1,))
    kontrol("API anahtarlı parça girmedi, aynı dosyanın diğer parçası girdi", "notlar/anahtar.md" in bul(c, "Rize lojistik") and not bul(c, "sk-ant-api03") and not bul(c, "ödeme servisi hesap bağlantısı"))
    kontrol("kart ve IBAN parçası girmedi, 'şifreleme politikası' girdi", not bul(c, "4539") and not bul(c, "IBAN hesabına") and "notlar/odeme.md" in bul(c, "şifreleme politikası"))
    kontrol("adla hariç ve kopya: çevrilmedi, dizinde yok", not c.execute("SELECT 1 FROM dosya WHERE yol IN ('belgeler/kimlik-on.jpg', 'eski/depo-sayimi.png')").fetchone())
    on = [os.path.join(d, f) for d, _, fs in os.walk(m.ONBELLEK) for f in fs]
    kontrol("önbellek uygulama klasöründe, 0600, proje klasörüne yazılmadı", len(on) == 4 and all((os.stat(f).st_mode & 0o777) == 0o600 for f in on)
            and not any("cevirme" in f or f.endswith(".json") for _, _, fs in os.walk(P) for f in fs))
    r2 = m.cevir(c)
    kontrol("ikinci tur: hepsi önbellekte, 0 çevirme", r2["cevrildi"] == 0 and r2["onbellekte"] == 4, json.dumps(r2)[:200])
    # zaman aşımı: takılan işçi → dosya zaman_asimi, sonraki dosya yeni işçiyle
    sahte = os.path.join(T, "sahte-isci.py")
    open(sahte, "w").write("import sys, json, time\nprint(json.dumps({'hazir': True, 'arac': 'sahte'}), flush=True)\nfor l in sys.stdin:\n"
                           "    o = json.loads(l)\n    if 'takil' in o['yol']: time.sleep(30)\n    print(json.dumps({'id': o['id'], 'parcalar': [['', 'sahte metin ' + o['yol'][-12:]]]}), flush=True)\n")
    shutil.copy(os.path.join(P, "ekran/depo-sayimi.png"), os.path.join(P, "ekran/takil.png")); open(os.path.join(P, "ekran/takil.png"), "ab").write(b"1")
    shutil.copy(os.path.join(P, "ekran/depo-sayimi.png"), os.path.join(P, "ekran/sonraki.png")); open(os.path.join(P, "ekran/sonraki.png"), "ab").write(b"2")
    os.utime(os.path.join(P, "ekran/takil.png"), (time.time() + 5, time.time() + 5))  # en yeni önce çevrilir
    m.CEVIR_ISCI = sahte; m.AYAR["cevir_sure_sn"] = 2; t0 = time.time(); r3 = m.cevir(c)
    kontrol("takılan dosya zaman aşımı, sonraki dosya yeni işçiyle çevrildi", r3["hatalar"] == [("ekran/takil.png", "zaman_asimi")] and r3["cevrildi"] == 1
            and time.time() - t0 < 20 and "ekran/sonraki.png" in bul(c, "sahte metin"), json.dumps(r3)[:300])
    kontrol("zaman aşımı yeniden denenmez; --yeniden-dene ile denenir", m.cevir(c)["hata"] == 0 and m.cevir(c, yeniden_dene=True)["hatalar"] == [("ekran/takil.png", "zaman_asimi")])
    # toplantı: dosyalar arasında durur, sonra kaldığı yerden
    for i in range(3):
        y = os.path.join(P, f"ekran/yeni{i}.png"); shutil.copy(os.path.join(P, "ekran/depo-sayimi.png"), y); open(y, "ab").write(f"yeni{i}".encode())
    json.dump({"meeting": {"title": "Deneme"}}, open(os.path.join(U, "canli", "heartbeat.json"), "w"))
    r4 = m.cevir(c)
    kontrol("toplantı sürerken çevirmedi, nedeni yazdı", r4["cevrildi"] == 0 and "toplantı" in (r4["durdu"] or ""), json.dumps(r4)[:200])
    os.remove(os.path.join(U, "canli", "heartbeat.json")); r5 = m.cevir(c)
    kontrol("toplantı bitince kaldığı yerden: yalnız yeni 3 dosya", r5["cevrildi"] == 3 and r5["kalan"] == 0, json.dumps(r5)[:200])
    # çevrilmiş metin sürüm benzerliğine girer: tarihli adlı iki taranmış PDF, aynı metin → sürüm; farklı metin → değil
    m.CEVIR_ISCI = os.path.join(KOD, "cevir-isci.py"); m.AYAR.pop("cevir_sure_sn", None)
    URET2 = "import sys\nfrom PIL import Image, ImageDraw, ImageFont\nf = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 30)\n" \
            "for yol, metin in zip(sys.argv[1::2], sys.argv[2::2]):\n    im = Image.new('RGB', (1500, 700), 'white'); d = ImageDraw.Draw(im)\n" \
            "    for i, s in enumerate(metin.split('|')): d.text((50, 50 + 60 * i), s, fill='black', font=f)\n    im.save(yol)\n"
    ORT_M = "Haftalık depo raporu: Bursa deposunda sayım tamamlandı|eksik ürün yok, iade kutuları ayrıldı|sevkiyat planı perşembe günü güncellenecek|forklift bakımı ay sonunda yapılacak|personel vardiyası değişmedi"
    os.makedirs(f"{P}/rapor", exist_ok=True)
    subprocess.run([VPY, "-B", "-c", URET2, f"{P}/rapor/depo-20261016.pdf", ORT_M, f"{P}/rapor/depo-20261017.pdf", ORT_M + "|ek not: raf etiketleri yenilendi",
                    f"{P}/rapor/fatura-20261016.pdf", "Fatura özeti: Konya bayisi|tutar ve vade bilgisi|ödeme koşulları otuz gün", f"{P}/rapor/fatura-20261017.pdf",
                    "Toplantı tutanağı: Samsun ziyareti|katılımcılar ve kararlar|sonraki adımlar listesi"], check=True, capture_output=True)
    for k, g in (("rapor/depo-20261016.pdf", -60), ("rapor/depo-20261017.pdf", -30), ("rapor/fatura-20261016.pdf", -60), ("rapor/fatura-20261017.pdf", -30)):
        os.utime(os.path.join(P, k), (time.time() + g, time.time() + g))
    m.cevir(c); rt = m.tara(c); dd = {x["yol"]: x for x in rt["dosyalar"]}
    kontrol("çevrilmiş PDF: aynı metinli tarihli adlar sürüm, farklı metinliler değil",
            dd["rapor/depo-20261016.pdf"]["gecerli"] == 0 and dd["rapor/depo-20261017.pdf"]["gecerli"] == 1
            and dd["rapor/fatura-20261016.pdf"]["surum_grubu"] is None and dd["rapor/fatura-20261017.pdf"]["surum_grubu"] is None, json.dumps(rt["surum_gruplari"], ensure_ascii=False)[:300])
    kontrol("eski sürüm PDF arama sonucunda yok", [x[1] for x in m.ara("forklift bakımı", n=5, guncel=False)[0]] == ["rapor/depo-20261017.pdf"])
    open(os.path.join(P, "ekran/butce.png"), "wb").write(open(os.path.join(P, "ekran/depo-sayimi.png"), "rb").read() + b"butce")
    rb = m.cevir(c, butce_sn=1e-6)
    kontrol("tur süresi dolunca durur, dosya sonraki tura kalır", rb["cevrildi"] == 0 and "süre" in (rb["durdu"] or "") and rb["kalan"] == 1, json.dumps(rb)[:200])
    os.remove(os.path.join(P, "ekran/butce.png"))
    # v0.121.0: Pages, Numbers, .xls — kurgu örnekler depo dışında (../_yedek/test-araclari/bicim/; Pages/Numbers'la ve xlwt ile üretildi); yoksa atlanır
    BICIM = os.path.join(os.path.dirname(KOD), "_yedek", "test-araclari", "bicim")
    if os.path.isdir(BICIM) and subprocess.run([VPY, "-c", "import xlrd"], capture_output=True).returncode == 0:
        os.makedirs(f"{P}/bicim", exist_ok=True)
        for f in ("ornek.pages", "ornek.numbers", "ornek.xls"): shutil.copy(os.path.join(BICIM, f), f"{P}/bicim/{f}")
        rb2 = m.cevir(c)
        kontrol("Pages, Numbers, .xls çevrildi", rb2["cevrildi"] == 3 and rb2["hata"] == 0, json.dumps(rb2, ensure_ascii=False)[:300])
        kontrol("Pages metni aranıyor; .xls ve Numbers satırları aranıyor", "bicim/ornek.pages" in bul(c, "depo taşınmasını İzmir")
                and {"bicim/ornek.xls", "bicim/ornek.numbers"} <= bul(c, "Elektrik 830"), str(bul(c, "Elektrik 830")))
        import sqlite3
        v = sqlite3.connect(m.PROJE_VT); tb = dict(v.execute("SELECT dosya, tablo FROM tablolar WHERE dosya LIKE 'bicim/%'").fetchall())
        top = {d: v.execute(f'SELECT sum(tutar) FROM "{t}"').fetchone()[0] for d, t in tb.items()}
        tr = v.execute(f'SELECT tarih FROM "{tb.get("bicim/ornek.xls", "yok")}" ORDER BY satir').fetchall() if "bicim/ornek.xls" in tb else []; v.close()
        kontrol("Numbers ve .xls sayfaları proje veritabanında: toplam 13330.5, .xls tarihi ISO", top == {"bicim/ornek.xls": 13330.5, "bicim/ornek.numbers": 13330.5}
                and tr == [("2026-11-01",), ("2026-11-05",)], f"{top} {tr}")
    else: print("  (Pages/Numbers/.xls örnekleri ya da xlrd yok — o kısım atlandı)")
    m.CEVIR_ISCI = os.path.join(T, "yok.py")
    open(os.path.join(P, "ekran/son.png"), "wb").write(open(os.path.join(P, "ekran/depo-sayimi.png"), "rb").read() + b"3")
    r6 = m.cevir(c)
    kontrol("işçi açılamazsa durur, sebep yazılır, dosya önbelleğe yazılmaz", r6["cevrildi"] == 0 and "işçi açılmadı" in (r6["durdu"] or "") and r6["kalan"] == 1, json.dumps(r6)[:200])
finally:
    c.close(); shutil.rmtree(T, ignore_errors=True)
print(f"çevirme: {len(hatalar)} sorun" if hatalar else "çevirme: ✓ hepsi geçti")
sys.exit(1 if hatalar else 0)
